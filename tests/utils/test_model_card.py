import ast
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from openpyxl import load_workbook

from yolox.model_card import comparison, dataset_snapshot, write_report
from yolox.model_report_html import render_html


class TestModelCard(unittest.TestCase):
    def test_local_base_is_shown_in_report(self):
        html = render_html({"model": {"base_version": "none", "base_source": "archivo local: /weights/yolox_s.pth"}})
        self.assertIn("archivo local: /weights/yolox_s.pth", html)
        self.assertNotIn("como primer modelo publicado", html)

    def test_html_curves_escape_and_missing_history(self):
        report = {"model": {"project": "<script>alert(1)</script>", "version": "1.0.1",
                            "base_version": "1.0.0"},
                  "best": {"epoch": 2, "map_50_95": .6}, "dataset": {},
                  "comparison": {"same_validation": False,
                                "metrics": {"map_50_95": {"base": .5, "new": .6, "difference": .1}}},
                  "evaluations": [{"epoch": 1, "map_50_95": .5},
                                  {"epoch": 2, "map_50_95": .6}],
                  "training_history": [{"epoch": 1, "total_loss": .8, "iou_loss": .3,
                                        "learning_rate": .001},
                                       {"epoch": 2, "total_loss": .4, "iou_loss": .2,
                                        "learning_rate": .0001}]}
        html = render_html(report)
        self.assertEqual(html.count('<svg '), 4)
        self.assertIn('época 2: 60.00%', html)
        self.assertIn('comparación es histórica', html)
        self.assertNotIn('<script>', html)
        self.assertIn('&lt;script&gt;', html)
        self.assertNotIn('https://', html)
        report.pop('training_history')
        self.assertEqual(render_html(report).count('No hay historial disponible'), 3)

    def test_no_base_comparison_skips_validation_warning(self):
        # A first-ever model (no base at all) must not warn about an unverified
        # comparison, since no comparison is actually being shown.
        report = {"model": {"project": "demo"}, "best": {"epoch": 1}, "dataset": {},
                  "comparison": {"same_validation": False, "metrics": {}},
                  "evaluations": [], "training_history": []}
        html = render_html(report)
        self.assertNotIn('comparación es histórica', html)
        self.assertIn('primer modelo', html)

    def test_training_losses_accept_tensors_and_disabled_l1_scalar(self):
        tree = ast.parse(Path("yolox/core/trainer.py").read_text())
        trainer = next(n for n in tree.body if isinstance(n, ast.ClassDef))
        method = next(n for n in trainer.body if isinstance(n, ast.FunctionDef)
                      and n.name == "record_training_losses")
        namespace = {}
        exec(compile(ast.Module(body=[method], type_ignores=[]), "trainer.py", "exec"), namespace)

        class Tensor:
            def __init__(self, value):
                self.value = value

            def detach(self):
                return self

            def item(self):
                return self.value

        fake = SimpleNamespace(card_loss_sums={}, card_loss_count=0)
        record = namespace["record_training_losses"]
        record(fake, {"total_loss": Tensor(.5), "l1_loss": 0.0, "num_fg": 4})
        record(fake, {"total_loss": Tensor(.25), "l1_loss": Tensor(.125)})
        self.assertEqual(fake.card_loss_sums, {"total_loss": .75, "l1_loss": .125})
        self.assertEqual(fake.card_loss_count, 2)

    def test_dataset_identity_and_test_alias(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "images").mkdir()
            (root / "images/a.jpg").write_bytes(b"image content")
            data = {"images": [{"id": 1, "file_name": "a.jpg"}],
                    "categories": [{"id": 3, "name": "control"}],
                    "annotations": [{"category_id": 3}]}
            (root / "train.json").write_text(json.dumps(data))
            exp = SimpleNamespace(data_dir=tmp, annotations_dir=".", train_ann="train.json",
                                  val_ann="train.json", test_ann="train.json", train_image_dir="images")
            snapshot = dataset_snapshot(exp)
            self.assertEqual(snapshot["total_images"], 1)
            self.assertIsNone(snapshot["splits"]["test"])
            self.assertEqual(snapshot["splits"]["train"]["annotations_per_class"], {"control": 1})
            self.assertTrue(snapshot["identity_complete"])

    def test_comparison_missing_and_growth(self):
        ds = {"splits": {}, "identity_complete": True, "image_hashes": ["a", "b"]}
        report = {"dataset": ds, "best": {"map_50_95": .6}}
        base = {"dataset": {**ds, "image_hashes": ["a"]}, "best": {"map_50_95": .5}}
        result = comparison(report, base)
        self.assertAlmostEqual(result["metrics"]["map_50_95"]["difference"], .1)
        self.assertIsNone(result["metrics"]["precision"]["difference"])
        self.assertEqual(result["dataset"], {"new_images": 1, "existing_images": 1, "growth_percent": 100})
        self.assertIsNone(comparison(report, None)["dataset"]["new_images"])

    def test_persistent_report_no_overwrite_and_best_selection(self):
        with tempfile.TemporaryDirectory() as tmp:
            report = {"model": {"version": "1.0.1"}, "dataset": {"splits": {}},
                      "best": {"map_50_95": .8, "epoch": 2},
                      "evaluations": [{"map_50_95": .8}, {"map_50_95": .3}]}
            write_report(report, tmp)
            saved = json.loads((Path(tmp) / "metrics.json").read_text())
            self.assertEqual(saved["comparison"]["metrics"]["map_50_95"]["new"], .8)
            self.assertIn("N/A", (Path(tmp) / "model_report.html").read_text(encoding="utf-8"))
            workbook = load_workbook(Path(tmp) / "model_report.xlsx")
            self.assertNotIn("Gráficas", workbook.sheetnames)
            self.assertTrue((Path(tmp) / "model_report.md").is_file())
            self.assertEqual(workbook["Formato"]["K10"].value, "1.0.1")
            with self.assertRaises(FileExistsError):
                write_report(report, tmp)

    def test_excel_report_uses_markdown_and_measured_metrics(self):
        with tempfile.TemporaryDirectory() as tmp:
            report = {"model": {"project": "Prueba", "name": "Detector", "version": "1.2.3"},
                      "dataset": {"splits": {"train": {"images": 8}, "val": {"images": 2},
                                             "test": {"images": 0}}, "total_images": 10},
                      "training": {"max_epoch": 2, "batch_size": 4},
                      "best": {"epoch": 2, "map_50_95": 0.6, "ap50": 0.8},
                      "evaluations": [{"epoch": 1, "map_50_95": 0.4, "ap50": 0.6},
                                      {"epoch": 2, "map_50_95": 0.6, "ap50": 0.8}],
                      "training_history": [{"epoch": 1, "total_loss": 2, "iou_loss": 1,
                                            "learning_rate": .01},
                                           {"epoch": 2, "total_loss": 1, "iou_loss": .5,
                                            "learning_rate": .001}]}
            write_report(report, tmp)
            workbook = load_workbook(Path(tmp) / "model_report.xlsx")
            self.assertEqual(workbook["Formato"]["C10"].value, "Prueba")
            self.assertEqual(workbook["Formato"]["G40"].value,
                             "train: 80.0% / val: 20.0% / test: 0.0%")
            self.assertEqual(workbook["Formato"]["D70"].value, 0.6)
            self.assertEqual(workbook["Formato"]["F70"].value, "90 %")
            self.assertIn("umbrales IoU", workbook["Formato"]["H70"].value)
            self.assertNotIn("Gráficas", workbook.sheetnames)
            self.assertEqual(workbook["Control de Cambios"]["C3"].value, "Versión inicial")

            markdown = Path(tmp) / "model_report.md"
            content = markdown.read_text(encoding="utf-8")
            markdown.write_text(content.replace("- cliente:", "- cliente: Clínica Veterinaria")
                                .replace("- map_50_95: Precisión", "- map_50_95: Calidad personalizada. Precisión"),
                                encoding="utf-8")
            from yolox.model_report_excel import write_excel_report
            report["model"]["experiment"] = "/tmp/entrenamiento.py"
            write_excel_report(report, Path(tmp) / "model_report.xlsx", markdown)
            workbook = load_workbook(Path(tmp) / "model_report.xlsx")
            self.assertEqual(workbook["Formato"]["K14"].value, "Clínica Veterinaria")
            self.assertIn("Calidad personalizada", workbook["Formato"]["H70"].value)
            self.assertEqual(workbook["Formato"]["G59"].value, "/tmp/entrenamiento.py")
            self.assertEqual(workbook["Formato"]["D70"].value, 0.6)

    def test_failed_training_does_not_generate_card(self):
        # Execute the actual orchestration method without requiring CUDA/PyTorch.
        tree = ast.parse(Path("yolox/core/trainer.py").read_text())
        trainer = next(n for n in tree.body if isinstance(n, ast.ClassDef))
        method = next(n for n in trainer.body if isinstance(n, ast.FunctionDef) and n.name == "train")
        namespace = {"logger": SimpleNamespace(error=lambda *a: None)}
        exec(compile(ast.Module(body=[method], type_ignores=[]), "trainer.py", "exec"), namespace)
        calls = []
        fake = SimpleNamespace(rank=0, before_train=lambda: None,
                               train_in_epoch=lambda: None,
                               write_model_card=lambda: calls.append("card"),
                               after_train=lambda: calls.append("cleanup"))
        namespace["train"](fake)
        self.assertEqual(calls, ["card", "cleanup"])
        calls.clear()
        def fail():
            raise RuntimeError("training failed")
        fake.train_in_epoch = fail
        with self.assertRaises(RuntimeError):
            namespace["train"](fake)
        self.assertEqual(calls, ["cleanup"])


if __name__ == "__main__":
    unittest.main()
